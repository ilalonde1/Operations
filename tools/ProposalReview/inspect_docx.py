import sys, re, zipfile
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document

path = sys.argv[1]
d = Document(path)

print("=== BODY PARAGRAPHS (index | style | text[:90]) — non-empty only ===")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        print(f"{i:4d} | {p.style.name[:22]:22s} | {t[:90]}")

print("\n=== TABLES (index | rows x cols | row0 cells) ===")
for ti, t in enumerate(d.tables):
    r0 = [c.text.strip().replace('\n', ' / ')[:28] for c in t.rows[0].cells]
    print(f"T{ti} | {len(t.rows)}x{len(t.columns)} | {r0}")

print("\n=== FEE TABLE full dump (the table whose row0 has 'Personnel') ===")
for ti, t in enumerate(d.tables):
    if any('Personnel' in c.text for c in t.rows[0].cells):
        for ri, row in enumerate(t.rows):
            cells = [c.text.strip().replace('\n', ' / ') for c in row.cells]
            print(f"  T{ti} r{ri}: {cells}")
        # runs per cell in row 1 (to see fragmentation)
        for ci, c in enumerate(t.rows[1].cells):
            runs = [r.text for p in c.paragraphs for r in p.runs]
            print(f"    r1 c{ci} runs={runs}")

print("\n=== RAW XML search: text boxes / headers for cover strings ===")
z = zipfile.ZipFile(path)
for name in z.namelist():
    if name.startswith('word/') and name.endswith('.xml'):
        x = z.read(name).decode('utf-8', 'replace')
        for key in ['RFP 061', 'CUMMUNITY', 'Proposal for Structural', 'West Campus Recreation Center Phase 1', 'Sincerely', 'Overhead Rate Proposed', 'Other team members']:
            n = x.count(key)
            if n:
                print(f"  {name}: '{key}' x{n}" + ("  (in txbxContent)" if 'txbxContent' in x and key in x and x.find(key) > x.find('txbxContent') > -1 else ""))
